"""Audit trail, users, bulk import, reports, live-feed control and demo reset."""
import csv
import io

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import Response

from .. import audit, planning, state, storage
from ..db import iso, utcnow
from ..live import hub
from ..models import OTP_ROLES, Aircraft, AuditLog, Engine, Setting, Squadron, User
from ..schemas import UserIn, UserPatch
from ..security import Ctx, fail, hash_password, need
from .common import finish, squadron_for

router = APIRouter(tags=["admin"])
AUDIT_ROLES = ("AUDITOR", "CO", "ADMIN", "HQ")


def setting(db, key, default=""):
    s = db.get(Setting, key)
    return s.value if s else default


def put_setting(db, key, value):
    s = db.get(Setting, key)
    if s:
        s.value = value
    else:
        db.add(Setting(key=key, value=value))


# ---------- audit ----------

@router.get("/audit")
def audit_trail(limit: int = 500, ctx: Ctx = Depends(need(*AUDIT_ROLES))):
    rows = ctx.db.query(AuditLog).order_by(AuditLog.id.desc()).limit(min(limit, 2000)).all()
    return [{"id": e.id, "ts": e.ts, "actor": e.actor, "role": e.role, "action": e.action, "detail": e.detail,
             "prev_hash": e.prev_hash, "hash": e.hash} for e in rows]


@router.get("/audit/verify")
def audit_verify(ctx: Ctx = Depends(need(*AUDIT_ROLES))):
    return audit.verify(ctx.db)


# ---------- users (Admin only) ----------

def user_dict(u: User) -> dict:
    return {"id": u.id, "email": u.email, "name": u.name, "role": u.role, "squadron_code": u.squadron_code,
            "otp_required": u.otp_required, "active": u.active, "failed_logins": u.failed_logins,
            "locked": bool(u.locked_until and u.locked_until > utcnow())}


@router.get("/users")
def users(ctx: Ctx = Depends(need("ADMIN"))):
    return [user_dict(u) for u in ctx.db.query(User).order_by(User.id)]


def _check_squadron(db, role, squadron_code):
    if role in ("CO", "ENGO", "TECH") and not squadron_code:
        fail(422, f"Role {role} must belong to a squadron.")
    if squadron_code and not db.get(Squadron, squadron_code):
        fail(404, f"Squadron {squadron_code} not found.")


@router.post("/users")
def add_user(body: UserIn, ctx: Ctx = Depends(need("ADMIN"))):
    email = body.email.lower()
    if ctx.db.query(User).filter(User.email == email).first():
        fail(409, f"A user with the email {email} already exists.")
    _check_squadron(ctx.db, body.role, body.squadron_code)
    u = User(email=email, name=body.name, role=body.role, squadron_code=body.squadron_code,
             password_hash=hash_password(body.password), otp_required=body.role in OTP_ROLES)
    ctx.db.add(u)
    ctx.db.flush()
    return finish(ctx, "USER_ADDED", f"User {email} added with role {body.role}", user=user_dict(u))


@router.patch("/users/{user_id}")
def edit_user(user_id: int, body: UserPatch, ctx: Ctx = Depends(need("ADMIN"))):
    u = ctx.db.get(User, user_id)
    if not u:
        fail(404, f"User {user_id} not found.")
    if u.id == ctx.user.id and body.active is False:
        fail(409, "You cannot deactivate your own account.")
    changes = body.model_dump(exclude_none=True)
    _check_squadron(ctx.db, changes.get("role", u.role), changes.get("squadron_code", u.squadron_code))
    if changes.pop("unlock", None):
        u.failed_logins, u.locked_until = 0, None
    if "password" in changes:
        u.password_hash = hash_password(changes.pop("password"))
    for k, v in changes.items():
        setattr(u, k, v)
    u.otp_required = u.role in OTP_ROLES
    ctx.db.flush()
    return finish(ctx, "USER_EDITED", f"User {u.email} updated ({', '.join(body.model_dump(exclude_none=True).keys())})",
                  message=f"User {u.email} updated.", user=user_dict(u))


# ---------- bulk import of technical records ----------

IMPORT_COLUMNS = ["tail_no", "flying_hours", "engine_position", "hours_since_overhaul"]


def _read_rows(filename: str, data: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip().lower() if h is not None else "" for h in rows[0]] if rows else []
        return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        return [{(k or "").strip().lower(): v for k, v in row.items()} for row in reader]
    fail(422, "Please upload a .csv or .xlsx file.")


@router.post("/import/records")
async def import_records(file: UploadFile = File(...), ctx: Ctx = Depends(need("ADMIN", "ENGO"))):
    """Columns: tail_no, flying_hours, engine_position (1 or 2), hours_since_overhaul. Good rows are saved; each bad row gets a clear message."""
    data = await file.read()
    try:
        rows = _read_rows(file.filename or "", data)
    except (UnicodeDecodeError, KeyError, ValueError, OSError):
        fail(422, "The file could not be read. Save it as CSV (UTF-8) or Excel .xlsx and try again.")
    if not rows:
        fail(422, "The file has no data rows.")
    missing = [c for c in IMPORT_COLUMNS if c not in rows[0]]
    if missing:
        fail(422, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(IMPORT_COLUMNS)}.")
    errors, saved, tails = [], 0, set()
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        tail = str(row.get("tail_no") or "").strip()
        ac = ctx.db.get(Aircraft, tail) if tail else None
        if not ac or (ctx.user.role == "ENGO" and ac.squadron_code != ctx.squadron):
            errors.append({"row": n, "message": f"Row {n}: aircraft '{tail}' not found." if tail else f"Row {n}: tail_no is empty."})
            continue
        try:
            fh, hso, pos = float(row["flying_hours"]), float(row["hours_since_overhaul"]), int(float(row["engine_position"]))
        except (TypeError, ValueError):
            errors.append({"row": n, "message": f"Row {n}: flying_hours, engine_position and hours_since_overhaul must be numbers."})
            continue
        if pos not in (1, 2):
            errors.append({"row": n, "message": f"Row {n}: engine_position must be 1 or 2 (got {pos})."})
            continue
        if fh < 0 or hso < 0:
            errors.append({"row": n, "message": f"Row {n}: hours cannot be negative."})
            continue
        if fh < ac.flying_hours:
            errors.append({"row": n, "message": f"Row {n}: flying_hours {fh:g} is lower than the current record ({ac.flying_hours:g}). Hours cannot go down."})
            continue
        eng = ctx.db.query(Engine).filter(Engine.tail_no == tail, Engine.position == pos).first()
        if not eng:
            errors.append({"row": n, "message": f"Row {n}: {tail} has no engine in position {pos}."})
            continue
        ac.flying_hours, eng.hours_since_overhaul = fh, hso
        saved += 1
        tails.add(tail)
    ctx.db.flush()
    stored = storage.save(f"imports/{utcnow():%Y%m%dT%H%M%S}-{file.filename}", data)
    out = finish(ctx, "RECORDS_IMPORTED", f"Technical records import '{file.filename}': {saved} rows saved, {len(errors)} rows rejected",
                 list(tails), message=f"{saved} row(s) saved, {len(errors)} row(s) rejected.")
    from .. import jobs
    jobs.overhaul_check(ctx.db)
    return {**out, "saved": saved, "rejected": len(errors), "errors": errors, "stored_file": stored}


@router.get("/files/link")
def file_link(name: str, ctx: Ctx = Depends(need("ADMIN", "AUDITOR"))):
    url = storage.signed_link(name)
    if not url:
        fail(404, "File storage is not configured or the file does not exist.")
    return {"url": url, "expires_minutes": 10}


# ---------- reports ----------

def _report_rows(db, sq):
    score = state.fleet_score(db, sq)
    fc = planning.dashboard_forecast(db, sq)
    return score, fc, [state.aircraft_dict(a) for a in state.squadron_aircraft(db, sq)]


@router.get("/reports/readiness.csv")
def readiness_csv(squadron: str | None = None, ctx: Ctx = Depends(need("CO", "ENGO", "FSO", "AUDITOR", "HQ"))):
    sq = squadron_for(ctx, squadron)
    score, fc, acs = _report_rows(ctx.db, sq)
    label = setting(ctx.db, "classification", "DEMO DATA – UNCLASSIFIED")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label])
    w.writerow([f"Readiness report {sq}", f"generated {iso(utcnow())}"])
    w.writerow(["Fleet Health Score", score["score"]])
    for r in score["rules"]:
        w.writerow([r["label"], r["points"], r["detail"]])
    w.writerow([])
    w.writerow(["tail_no", "status", "readiness", "health_score", "flying_hours"])
    for a in acs:
        w.writerow([a["tail_no"], a["status_label"], a["readiness"], a["health_score"], a["flying_hours"]])
    w.writerow([])
    w.writerow(["day", "mission_capable_without_plan", "mission_capable_with_plan"])
    for d in fc["days"]:
        w.writerow([d, fc["without"][d], fc["with_plan"][d]])
    w.writerow([label])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="readiness-{sq}.csv"'})


@router.get("/reports/readiness.pdf")
def readiness_pdf(squadron: str | None = None, ctx: Ctx = Depends(need("CO", "ENGO", "FSO", "AUDITOR", "HQ"))):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    sq = squadron_for(ctx, squadron)
    score, fc, acs = _report_rows(ctx.db, sq)
    label = setting(ctx.db, "classification", "DEMO DATA – UNCLASSIFIED").replace("–", "-")
    styles = getSampleStyleSheet()
    buf = io.BytesIO()

    def banner(canvas, doc):  # classification banner at the top and bottom of every page
        canvas.setFont("Helvetica-Bold", 10)
        canvas.setFillColor(colors.HexColor("#b45309"))
        canvas.drawCentredString(A4[0] / 2, A4[1] - 28, label)
        canvas.drawCentredString(A4[0] / 2, 20, label)

    k = score["kpis"]
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f2545")),
                       ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8)])
    story = [
        Paragraph(f"VAYU-READY readiness report: {score['squadron_name']} ({sq})", styles["Title"]),
        Paragraph(f"Generated {iso(utcnow())}. Advisory. Final decision rests with the authorised engineer.", styles["Normal"]),
        Spacer(1, 10), Paragraph(f"Fleet Health Score: {score['score']} / 100", styles["Heading2"]),
        Table([["Rule", "Points", "Why"]] + [[r["label"], r["points"], r["detail"]] for r in score["rules"]], style=grid,
              colWidths=[190, 40, 290]),
        Spacer(1, 10),
        Paragraph(f"Mission-capable {k['mission_capable']}/{k['total']}; forecast in {k['forecast_day']} days {k['forecast_mission_capable']}/{k['total']}; "
                  f"on ground {k['aog']} ({k['aog_awaiting_spares']} awaiting spares); combat-ready {k['combat_ready']}; "
                  f"transport-ready {k['transport_ready']}. 30-day forecast: {fc['without'][-1]} without plan, {fc['with_plan'][-1]} with plan.", styles["Normal"]),
        Spacer(1, 10),
        Table([["Tail no.", "Status", "Readiness", "Health", "Flying hours"]] +
              [[a["tail_no"], a["status_label"], a["readiness"], a["health_score"], a["flying_hours"]] for a in acs], style=grid),
    ]
    SimpleDocTemplate(buf, pagesize=A4, topMargin=50, bottomMargin=40).build(story, onFirstPage=banner, onLaterPages=banner)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="readiness-{sq}.pdf"'})


# ---------- live feed ----------

@router.get("/stream/status")
def stream_status(ctx: Ctx = Depends(need())):
    return {"running": setting(ctx.db, "stream_running", "0") == "1", "broker_connected": hub.mqtt_connected,
            "messages_received": hub.messages, "last_message_at": hub.last_message_at, "dashboards_connected": len(hub.clients)}


@router.post("/stream/start")
def stream_start(ctx: Ctx = Depends(need("ADMIN"))):
    put_setting(ctx.db, "stream_running", "1")
    return finish(ctx, "STREAM_STARTED", "MQTT sensor replay started", message="Live sensor replay started.")


@router.post("/stream/stop")
def stream_stop(ctx: Ctx = Depends(need("ADMIN"))):
    put_setting(ctx.db, "stream_running", "0")
    return finish(ctx, "STREAM_STOPPED", "MQTT sensor replay stopped", message="Live sensor replay stopped.")


# ---------- demo mode ----------

@router.post("/demo/reset")
def demo_reset(ctx: Ctx = Depends(need("ADMIN"))):
    """Reloads the seed data so the demo always starts at score 71."""
    import seed
    actor, role = ctx.user.email, ctx.user.role
    ctx.db.close()
    score = seed.run(quiet=True)
    from ..db import engine
    engine.dispose()  # drop pooled connections that saw the old tables
    hub.refresh([], "DEMO_RESET")
    return {"message": f"Demo data reloaded. SQ7 Fleet Health Score is {score['score']}.", "score": score["score"], "by": actor, "role": role}
